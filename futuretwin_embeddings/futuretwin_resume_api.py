"""Privacy-conscious resume extraction and career-profile endpoints."""
from __future__ import annotations

import hashlib
import hmac
import csv
import io
import json
import os
import secrets
import shutil
import subprocess
import tempfile
import time
import unicodedata
import uuid
import zipfile
from datetime import timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any

import psycopg
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from docx import Document
from flask import jsonify, request, session
from futuretwin_local_extraction import extract_local_skills
from pypdf import PdfReader
from werkzeug.exceptions import RequestEntityTooLarge
from futuretwin_skills_profile import (
    MAX_ITEM_VALUE, ProfileValidationError,
    candidates_for, clean_user_profile, enrich_skills, merge_stored_profile,
    public_profile, validated_extraction, append_confirmed_skills,
)


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_RESUME_CHARS = 30_000
MAX_PDF_PAGES = 30
MAX_DOCX_UNCOMPRESSED = 35 * 1024 * 1024
PROFILE_TABLE = "futuretwin_career_profiles"


class ResumeRequestError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def _response(payload: dict[str, Any], status: int = 200):
    response = jsonify(payload)
    response.status_code = status
    response.headers["Cache-Control"] = "no-store, private"
    response.headers["Pragma"] = "no-cache"
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


def _session_secret() -> bytes:
    secret = os.getenv("SESSION_SECRET", "")
    if not secret:
        raise RuntimeError("Secure profile session configuration is missing.")
    return secret.encode("utf-8")


def _owner_hash(create: bool = False) -> bytes | None:
    token = session.get("ft_profile_owner")
    if not token and create:
        token = secrets.token_urlsafe(32)
        session["ft_profile_owner"] = token
        session.permanent = True
    if not token:
        return None
    return hmac.new(_session_secret(), token.encode("utf-8"), hashlib.sha256).digest()


def _csrf_token(create: bool = False) -> str | None:
    token = session.get("ft_profile_csrf")
    if not token and create:
        token = secrets.token_urlsafe(32)
        session["ft_profile_csrf"] = token
        session.permanent = True
    return token


def _csrf_valid() -> bool:
    expected = _csrf_token()
    provided = request.headers.get("X-CSRF-Token", "")
    return bool(expected and provided and hmac.compare_digest(expected, provided))


def _encryption_key() -> bytes:
    secret = _session_secret()
    return hashlib.sha256(b"futuretwin-career-profile-v1\0" + secret).digest()


def _encrypt_profile(profile: dict[str, Any]) -> bytes:
    nonce = secrets.token_bytes(12)
    ciphertext = AESGCM(_encryption_key()).encrypt(
        nonce,
        json.dumps(profile, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
        b"futuretwin-career-profile-v1",
    )
    return b"FT1" + nonce + ciphertext


def _decrypt_profile(blob: bytes) -> dict[str, Any]:
    if not blob or not blob.startswith(b"FT1") or len(blob) < 32:
        raise ValueError("Unsupported encrypted profile format.")
    plain = AESGCM(_encryption_key()).decrypt(
        blob[3:15], blob[15:], b"futuretwin-career-profile-v1"
    )
    data = json.loads(plain.decode("utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Stored profile data is invalid.")
    return data


def _database_connection():
    database_url = os.getenv("DATABASE_URL", "")
    if not database_url:
        raise RuntimeError("Profile storage is not configured.")
    return psycopg.connect(database_url, connect_timeout=6)


def _load_profile(owner: bytes) -> dict[str, Any] | None:
    with _database_connection() as conn:
        row = conn.execute(
            f"SELECT encrypted_payload, created_at, updated_at "
            f"FROM {PROFILE_TABLE} WHERE owner_hash = %s",
            (owner,),
        ).fetchone()
    if not row:
        return None
    profile = public_profile(_decrypt_profile(bytes(row[0])))
    profile["created_at"] = row[1].isoformat() if row[1] else None
    profile["updated_at"] = row[2].isoformat() if row[2] else None
    return profile


def _load_profile_document(owner: bytes) -> dict[str, Any] | None:
    """Load the encrypted owner document for server-side My Future operations."""
    with _database_connection() as conn:
        row = conn.execute(
            f"SELECT encrypted_payload FROM {PROFILE_TABLE} WHERE owner_hash = %s",
            (owner,),
        ).fetchone()
    return _decrypt_profile(bytes(row[0])) if row else None


def _save_profile(owner: bytes, profile: dict[str, Any], *, append: bool = False,
                  confirm_existing: bool = False) -> dict[str, Any]:
    profile_id = str(uuid.uuid4())
    with _database_connection() as conn:
        # Also serialize the very first save, when no row exists to lock yet.
        conn.execute("SELECT pg_advisory_xact_lock(%s)", (int.from_bytes(owner[:8], "big", signed=True),))
        # Preserve existing private legacy fields inside encryption. Do not trust
        # the client to return them, and serialize saves to an existing profile.
        previous_row = conn.execute(
            f"SELECT encrypted_payload FROM {PROFILE_TABLE} WHERE owner_hash = %s FOR UPDATE",
            (owner,),
        ).fetchone()
        previous = _decrypt_profile(bytes(previous_row[0])) if previous_row else {}
        added = duplicates = 0
        if append:
            stored, added, duplicates = append_confirmed_skills(
                previous, profile, confirm_existing=confirm_existing
            )
        else:
            stored = merge_stored_profile(previous, profile)
        encrypted = _encrypt_profile(stored)
        row = conn.execute(
            f"""
            INSERT INTO {PROFILE_TABLE}
                (profile_id, owner_hash, encrypted_payload)
            VALUES (%s, %s, %s)
            ON CONFLICT (owner_hash) DO UPDATE
            SET encrypted_payload = EXCLUDED.encrypted_payload,
                updated_at = NOW()
            RETURNING created_at, updated_at
            """,
            (profile_id, owner, encrypted),
        ).fetchone()
    result = public_profile(stored)
    result["created_at"] = row[0].isoformat() if row and row[0] else None
    result["updated_at"] = row[1].isoformat() if row and row[1] else None
    if append:
        result["_added_count"] = added
        result["_duplicate_count"] = duplicates
    return result


def _save_my_future_state(owner: bytes, state: dict[str, Any]) -> None:
    """Persist My Future state in the existing encrypted, session-owned profile."""
    profile_id = str(uuid.uuid4())
    with _database_connection() as conn:
        conn.execute(
            "SELECT pg_advisory_xact_lock(%s)",
            (int.from_bytes(owner[:8], "big", signed=True),),
        )
        previous_row = conn.execute(
            f"SELECT encrypted_payload FROM {PROFILE_TABLE} "
            "WHERE owner_hash = %s FOR UPDATE",
            (owner,),
        ).fetchone()
        if previous_row:
            stored = _decrypt_profile(bytes(previous_row[0]))
        else:
            stored = {"schema_version": 2, "skills": []}
        # Preserve every existing private field and the confirmed ESCO URIs.
        # Purpose is updated only by its field-scoped endpoint. Legacy/stale
        # full-state clients must not erase priorities saved in another tab.
        from futuretwin_my_future import read_state
        state["purpose"] = read_state(stored.get("my_future"))["purpose"]
        stored["my_future"] = state
        encrypted = _encrypt_profile(stored)
        conn.execute(
            f"""
            INSERT INTO {PROFILE_TABLE}
                (profile_id, owner_hash, encrypted_payload)
            VALUES (%s, %s, %s)
            ON CONFLICT (owner_hash) DO UPDATE
            SET encrypted_payload = EXCLUDED.encrypted_payload,
                updated_at = NOW()
            """,
            (profile_id, owner, encrypted),
        )


def _delete_profile(owner: bytes | None) -> bool:
    if owner is None:
        return False
    with _database_connection() as conn:
        cursor = conn.execute(
            f"DELETE FROM {PROFILE_TABLE} WHERE owner_hash = %s",
            (owner,),
        )
        return cursor.rowcount > 0


def _extract_text(file_bytes: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.casefold()
    if not file_bytes:
        raise ResumeRequestError("The selected document is empty.")

    try:
        if suffix == ".pdf":
            if not file_bytes.startswith(b"%PDF-"):
                raise ResumeRequestError("This file does not look like a valid PDF.")
            reader = PdfReader(io.BytesIO(file_bytes), strict=False)
            if reader.is_encrypted:
                raise ResumeRequestError("Password-protected PDFs are not supported.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ResumeRequestError("This PDF has too many pages. Try one with 30 pages or fewer.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                entries = archive.infolist()
                expanded_size = sum(entry.file_size for entry in entries)
                if (
                    len(entries) > 250
                    or expanded_size > MAX_DOCX_UNCOMPRESSED
                    or not any(entry.filename == "word/document.xml" for entry in entries)
                ):
                    raise ResumeRequestError("This Word document is not valid or is too large to process.")
            document = Document(io.BytesIO(file_bytes))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells if cell.text.strip()))
            text = "\n".join(parts)
        elif suffix == ".doc":
            if not file_bytes.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                raise ResumeRequestError("This file does not look like a valid Word document.")
            antiword = shutil.which("antiword")
            if not antiword:
                raise ResumeRequestError("Legacy .doc files cannot be read right now. Please use PDF or DOCX.")
            temporary_path = None
            try:
                with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as temporary:
                    temporary.write(file_bytes)
                    temporary.flush()
                    temporary_path = temporary.name
                parsed = subprocess.run(
                    [antiword, temporary_path],
                    check=False,
                    capture_output=True,
                    timeout=12,
                    env={"PATH": os.path.dirname(antiword), "HOME": tempfile.gettempdir()},
                )
                if parsed.returncode != 0:
                    raise ResumeRequestError("We couldn't read this Word document. Try saving it as DOCX.")
                text = parsed.stdout.decode("utf-8", errors="replace")
            except subprocess.TimeoutExpired as exc:
                raise ResumeRequestError("This Word document took too long to read. Try saving it as DOCX.") from exc
            finally:
                if temporary_path:
                    try:
                        os.remove(temporary_path)
                    except OSError:
                        pass
        else:
            raise ResumeRequestError("Choose a PDF or Word document (.doc or .docx).", 415)
    except ResumeRequestError:
        raise
    except Exception as exc:
        raise ResumeRequestError("We couldn't read this document. Check the file and try again.", 422) from exc

    text = unicodedata.normalize("NFKC", text)
    text = "".join(char for char in text if char in "\n\t" or ord(char) >= 32)
    text = text.strip()
    if not text:
        raise ResumeRequestError(
            "We couldn't find selectable text. Try a text-based PDF or Word document.",
            422,
        )
    if len(text) > MAX_RESUME_CHARS:
        raise ResumeRequestError(
            "This document is too long to process safely. Try a shorter resume.",
            413,
        )
    return text


def _extract_resume_skills(text: str) -> dict[str, Any]:
    return _extract_skills_local(text, "resume")


def _extract_skills_local(text: str, source_kind: str = "resume") -> dict[str, Any]:
    return extract_local_skills(text, source_kind)


def _validated_extraction(raw: dict[str, Any], source_text: str) -> dict[str, Any]:
    return validated_extraction(raw, source_text)


def _match_skill_candidates(skills: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from futuretwin_skill_api import match_skills

    return enrich_skills(skills, match_skills)


def _clean_user_profile(payload: Any) -> dict[str, Any]:
    from futuretwin_skill_api import match_skills
    try:
        return clean_user_profile(payload, match_skills)
    except ProfileValidationError as exc:
        raise ResumeRequestError(str(exc)) from exc


@lru_cache(maxsize=1)
def _passport_skill_catalog() -> dict[str, str]:
    path = Path(__file__).resolve().parent / "data/embedding_training/skill_catalog.csv"
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return {row["skillUri"].strip(): row["skill"].strip()
                for row in csv.DictReader(stream)
                if row.get("skillUri") and row.get("skill")}


def _catalog_skill(uri: str) -> dict[str, str] | None:
    """Resolve an exact local ESCO identity without running any model."""
    label = _passport_skill_catalog().get(uri)
    return {"uri": uri, "label": label} if label else None


def register_resume_routes(app) -> None:
    """Attach profile routes to the existing FutureTwin Flask application."""
    secret = os.getenv("SESSION_SECRET", "")
    if secret and not app.secret_key:
        app.secret_key = secret
    app.config.update(
        SESSION_COOKIE_NAME="futuretwin_profile_session",
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.getenv("REPLIT_DEPLOYMENT_ID") is not None,
        PERMANENT_SESSION_LIFETIME=timedelta(days=30),
        MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES + 256 * 1024,
    )
    from futuretwin_experience_api import register_experience_routes
    register_experience_routes(app)
    from futuretwin_intelligence_api import register_intelligence_routes
    register_intelligence_routes(app)

    @app.errorhandler(RequestEntityTooLarge)
    def api_profile_payload_too_large(_exc):
        return _response({"error": "The file must be 10 MB or smaller."}, 413)

    @app.get("/api/profile")
    def api_profile_get():
        try:
            owner = _owner_hash(create=True)
            profile = _load_profile(owner)
            return _response({"profile": profile, "csrf_token": _csrf_token(create=True)})
        except Exception:
            return _response({"error": "Your profile could not be loaded right now."}, 503)

    @app.post("/api/profile/extract")
    def api_profile_extract():
        try:
            if not _csrf_valid():
                return _response({"error": "Refresh the page and try again."}, 403)
            now = time.time()
            recent = [
                stamp for stamp in session.get("ft_resume_extract_times", [])
                if isinstance(stamp, (int, float)) and now - stamp < 3600
            ]
            if len(recent) >= 5:
                return _response({
                    "error": "Too many resume attempts. Please wait before trying again."
                }, 429)
            recent.append(now)
            session["ft_resume_extract_times"] = recent
            if request.form.get("consent", "").casefold() not in {"1", "true", "on", "yes"}:
                return _response({"error": "Please consent to AI processing before uploading."}, 400)
            uploaded = request.files.get("resume")
            if uploaded is None or not uploaded.filename:
                return _response({"error": "Choose a PDF or Word document to continue."}, 400)
            filename = Path(uploaded.filename).name
            if len(filename) > 180:
                return _response({"error": "This file name is too long."}, 400)
            data = uploaded.stream.read(MAX_UPLOAD_BYTES + 1)
            uploaded.stream.close()
            if len(data) > MAX_UPLOAD_BYTES:
                return _response({"error": "The file must be 10 MB or smaller."}, 413)
            source_text = _extract_text(data, filename)
            del data
            raw = _extract_resume_skills(source_text)
            draft = _validated_extraction(raw, source_text)
            draft["skills"] = _match_skill_candidates(draft["skills"])
            return _response({
                "draft": draft,
                "notice": "Review every item before saving. The original document is not stored.",
            })
        except ResumeRequestError as exc:
            return _response({"error": str(exc)}, exc.status)
        except Exception:
            return _response({
                "error": "We couldn't process this resume right now. Your document was not saved. Please try again later."
            }, 502)

    @app.post("/api/profile/match")
    def api_profile_match():
        try:
            if not _csrf_valid():
                return _response({"error": "Refresh the page and try again."}, 403)
            payload = request.get_json(silent=True)
            value = payload.get("value") if isinstance(payload, dict) else None
            if not isinstance(value, str) or not value.strip() or len(value) > MAX_ITEM_VALUE:
                raise ResumeRequestError("Enter a short skill or work activity.")
            from futuretwin_skill_api import match_skills
            return _response({"esco_candidates": candidates_for(value.strip(), match_skills)})
        except ResumeRequestError as exc:
            return _response({"error": str(exc)}, exc.status)
        except Exception:
            return _response({"error": "Skill matches could not be loaded. Please try again."}, 503)

    @app.post("/api/profile")
    def api_profile_save():
        try:
            if not _csrf_valid():
                return _response({"error": "Refresh the page and try again."}, 403)
            payload = request.get_json(silent=True)
            profile = _clean_user_profile(payload)
            owner = _owner_hash(create=True)
            saved = _save_profile(owner, profile)
            return _response({"profile": saved})
        except ResumeRequestError as exc:
            return _response({"error": str(exc)}, exc.status)
        except Exception:
            return _response({"error": "Your profile could not be saved right now. Please try again."}, 503)

    @app.post("/api/profile/skills/confirm")
    def api_profile_confirm_skill():
        try:
            if not _csrf_valid():
                return _response({"error": "Refresh the page and try again."}, 403)
            payload = request.get_json(silent=True)
            uri = payload.get("skill_uri") if isinstance(payload, dict) else None
            if not isinstance(uri, str) or not uri or len(uri) > 500:
                raise ResumeRequestError("Choose a valid ESCO skill.")
            if payload.get("confirmed") is not True:
                raise ResumeRequestError("Confirm that you have this skill before saving.")
            skill = _catalog_skill(uri)
            if not skill:
                raise ResumeRequestError("This ESCO skill is not available. Refresh and try again.")
            row = {
                "value": skill["label"], "evidence": "", "origin": "user_added",
                "confirmed": True, "selected_esco_uri": skill["uri"],
                "esco_candidates": [skill],
            }
            saved = _save_profile(
                _owner_hash(create=True), {"schema_version": 2, "skills": [row]},
                append=True, confirm_existing=True,
            )
            from futuretwin_my_future import _confirmed_skills
            return _response({
                "profile": saved,
                "confirmed_skills": list(_confirmed_skills(saved).values()),
                "added": saved.get("_added_count", 0) > 0,
            })
        except (ResumeRequestError, ProfileValidationError) as exc:
            return _response({"error": str(exc)}, getattr(exc, "status", 400))
        except Exception:
            return _response({"error": "The skill could not be saved. Please try again."}, 503)

    @app.delete("/api/profile")
    def api_profile_delete():
        try:
            if not _csrf_valid():
                return _response({"error": "Refresh the page and try again."}, 403)
            owner = _owner_hash(create=False)
            deleted = _delete_profile(owner)
            session.pop("ft_profile_owner", None)
            return _response({"deleted": deleted})
        except Exception:
            return _response({"error": "Your profile could not be deleted right now. Please try again."}, 503)

    @app.get("/api/my-future")
    def api_my_future_get():
        try:
            from futuretwin_my_future import (
                CAREER_AREAS, _confirmed_skills, future_scenario_framework,
                read_state, saved_careers,
            )

            owner = _owner_hash(create=True)
            stored = _load_profile_document(owner)
            state = read_state(stored.get("my_future") if stored else None)
            confirmed = _confirmed_skills(stored)
            return _response({
                "profile_exists": stored is not None,
                "confirmed_skill_count": len(confirmed),
                "confirmed_skills": list(confirmed.values()),
                "preferences": state["preferences"],
                "saved_occupations": state["saved_occupations"],
                "saved_careers": saved_careers(state),
                "roadmap_progress": state["roadmap_progress"],
                "roadmap_goals": state["roadmap_goals"],
                "insurance_alternatives": state["insurance_alternatives"],
                "purpose": state["purpose"],
                "career_areas": CAREER_AREAS,
                "future_scenarios": future_scenario_framework(),
                "csrf_token": _csrf_token(create=True),
            })
        except Exception:
            return _response({"error": "My Future could not be loaded right now."}, 503)

    @app.get("/api/my-future/recommendations")
    def api_my_future_recommendations():
        try:
            from futuretwin_my_future import build_recommendations, read_state

            try:
                limit = int(request.args.get("limit", "30"))
            except ValueError:
                return _response({"error": "limit must be a whole number."}, 400)
            if not 1 <= limit <= 30:
                return _response({"error": "limit must be from 1 to 30."}, 400)
            sort = request.args.get("sort", "best_match")
            if sort not in {"best_match", "lower_exposure", "growing_jobs", "purpose_alignment"}:
                return _response({"error": "Choose best_match, lower_exposure or growing_jobs."}, 400)
            owner = _owner_hash(create=True)
            stored = _load_profile_document(owner)
            state = read_state(stored.get("my_future") if stored else None)
            result = build_recommendations(stored, state, limit=limit, sort=sort)
            result["saved_occupations"] = state["saved_occupations"]
            result["preferences"] = state["preferences"]
            return _response(result)
        except Exception:
            return _response({
                "error": "Career recommendations could not be loaded right now."
            }, 503)

    @app.get("/api/my-future/career")
    def api_my_future_career():
        try:
            from futuretwin_my_future import build_career_details, read_state

            occupation_uri = request.args.get("uri", "").strip()
            if not occupation_uri or len(occupation_uri) > 500:
                return _response({"error": "Choose a career from ESCO search results."}, 400)
            owner = _owner_hash(create=True)
            stored = _load_profile_document(owner)
            state = read_state(stored.get("my_future") if stored else None)
            details = build_career_details(occupation_uri, stored, state)
            if details is None:
                return _response({"error": "ESCO occupation not found."}, 404)
            return _response(details)
        except Exception:
            return _response({
                "error": "Career details could not be loaded right now."
            }, 503)

    @app.post("/api/my-future")
    def api_my_future_save():
        try:
            if not _csrf_valid():
                return _response({"error": "Refresh the page and try again."}, 403)
            from futuretwin_my_future import (
                MyFutureValidationError, read_state, validate_state, saved_careers,
            )

            payload = request.get_json(silent=True)
            state = validate_state(payload)
            owner = _owner_hash(create=True)
            _save_my_future_state(owner, state)
            return _response({
                "preferences": state["preferences"],
                "saved_occupations": state["saved_occupations"],
                "saved_careers": saved_careers(state),
                "roadmap_progress": state["roadmap_progress"],
                "roadmap_goals": state["roadmap_goals"],
                "insurance_alternatives": state["insurance_alternatives"],
                "purpose": state["purpose"],
                "csrf_token": _csrf_token(create=True),
            })
        except MyFutureValidationError as exc:
            return _response({"error": str(exc)}, 400)
        except Exception:
            return _response({
                "error": "Your My Future changes could not be saved. Please try again."
            }, 503)


__all__ = ["register_resume_routes"]
